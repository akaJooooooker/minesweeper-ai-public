"""Build the English research manuscript using the Codex artifact runtime."""
from pathlib import Path
import json, math, re
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'output/research'
ASSETS=ROOT/'tmp/paper-qa/assets'
ASSETS.mkdir(parents=True,exist_ok=True)
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10,'axes.spines.top':False,'axes.spines.right':False,'savefig.dpi':220})

def wilson(w,n):
    p=w/n; z=1.959963984540054; d=1+z*z/n
    center=(p+z*z/(2*n))/d
    radius=z*math.sqrt(p*(1-p)/n+z*z/(4*n*n))/d
    return center-radius,center+radius

def exact(b,c):
    return min(1.,2*sum(math.comb(b+c,i) for i in range(min(b,c)+1))/2**(b+c))

result=json.loads((ROOT/'experiments/local-selfplay-20260906-result.json').read_text(encoding='utf-8'))
pairs=[json.loads(line) for line in (ROOT/'data/selfplay-v44-candidate-paired/pairs.jsonl').read_text(encoding='utf-8').splitlines()]
assert len(pairs)==5000 and len({(x['mode'],x['index']) for x in pairs})==5000
for label,key in [('reference','reference_wins'),('candidate','candidate_wins')]:
    assert sum(x[label]['won'] for x in pairs)==result[key]
assert exact(6,8)==result['mcnemar_exact_p']
for mode,expected in result['per_mode'].items():
    rows=[x for x in pairs if x['mode']==mode]
    assert len(rows)==expected['games']
    assert sum(x['reference']['won'] and not x['candidate']['won'] for x in rows)==expected['reference_only']
    assert sum(x['candidate']['won'] and not x['reference']['won'] for x in rows)==expected['candidate_only']

fig,ax=plt.subplots(figsize=(6.5,2.6))
for i,(label,values,color,offset) in enumerate([('V32',[1669,1371,657],'#777777',-.12),('V44',[1668,1388,690],'#176C84',.12)]):
    y=[v/20 for v in values]; ci=[wilson(v,2000) for v in values]
    ax.errorbar([j+offset for j in range(3)],y,yerr=[[y[j]-100*ci[j][0] for j in range(3)],[100*ci[j][1]-y[j] for j in range(3)]],fmt='o',capsize=4,color=color,label=label,markersize=5,linewidth=1.3)
ax.set_xticks(range(3),['Beginner','Intermediate','Expert'])
ax.set_ylabel('Win rate (%)'); ax.set_ylim(25,90); ax.set_xlim(-.6,2.6)
ax.grid(axis='y',color='#dddddd',lw=.6); ax.set_axisbelow(True)
ax.legend(frameon=False,loc='upper right',ncol=2)
fig.tight_layout(pad=.8); fig.savefig(ASSETS/'standard_rates.png'); plt.close(fig)

fig,ax=plt.subplots(figsize=(6.5,2.2)); ax.set_xlim(0,10); ax.set_ylim(0,3.1); ax.axis('off')
def box(x,y,w,h,text,fc='#f4f5f6'):
    ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=0.04,rounding_size=0.06',fc=fc,ec='#555555',lw=.8))
    ax.text(x+w/2,y+h/2,text,ha='center',va='center',fontsize=8.8,linespacing=1.4)
box(.04,.75,1.5,1.5,'Public board\n+ mine count')
box(1.98,.75,2.18,1.5,'Constraint inference\n24 to 40 cells\n40 to 64 if no guesses')
box(4.65,.35,3.15,2.3,'Proven safe: reveal\n\nExact: minimum risk\n\nNon-exact: neural risk\n+ gated top-3 policy')
box(8.25,.75,1.7,1.5,'Eligible action\nthrough adapter')
for a,b in [(1.55,1.94),(4.17,4.61),(7.82,8.2)]:
    ax.annotate('',xy=(b,1.5),xytext=(a,1.5),arrowprops=dict(arrowstyle='->',color='#333333',lw=1))
ax.text(5,.05,'Hidden mines and seed are not inference inputs',ha='center',fontsize=9,color='#333333')
fig.tight_layout(pad=.4); fig.savefig(ASSETS/'pipeline.png'); plt.close(fig)

source=(OUT/'minesweeper-study.md').read_text(encoding='utf-8')
source=re.sub(r'(?m)^(#{1,4} .+)\n(?!\n|#)',r'\1\n\n',source)
doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.2677); sec.page_height=Inches(11.6929)
sec.top_margin=Inches(.76); sec.bottom_margin=Inches(.76)
sec.left_margin=Inches(.83); sec.right_margin=Inches(.83)
sec.footer_distance=Inches(.32)
width=6.6077
styles=doc.styles
for style in styles:
    for el in list(style.element.iter()):
        if el.tag in (qn('w:pBdr'),qn('w:spacing')) and el.getparent() is not None and el.getparent().tag==qn('w:rPr'):
            el.getparent().remove(el)
        if el.tag==qn('w:pBdr') and el.getparent() is not None:
            el.getparent().remove(el)
        if el.tag==qn('w:rFonts'):
            for key in list(el.attrib):
                if 'Theme' in key or 'theme' in key: del el.attrib[key]
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    st=styles[name]; st.font.name='Times New Roman'; st.font.color.rgb=RGBColor(0,0,0)
    
    for family in ['ascii','hAnsi','eastAsia','cs']: st.element.get_or_add_rPr().rFonts.set(qn('w:'+family),'Times New Roman')
normal=styles['Normal']; normal.font.size=Pt(11)
normal.paragraph_format.line_spacing=1.10
normal.paragraph_format.space_after=Pt(6)
normal.paragraph_format.widow_control=True
for name,size,before,after in [('Heading 1',13,12,6),('Heading 2',11.5,9,4),('Heading 3',11,8,4)]:
    st=styles[name]; st.font.size=Pt(size); st.font.bold=True
    st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after)
    st.paragraph_format.keep_with_next=True
styles['Title'].font.size=Pt(23); styles['Title'].font.bold=True
styles['Title'].paragraph_format.space_after=Pt(8)
styles['Subtitle'].font.size=Pt(13); styles['Subtitle'].font.italic=False
styles['Subtitle'].paragraph_format.space_after=Pt(10)
styles['Caption'].font.bold=False; styles['Caption'].font.size=Pt(9.5); styles['Caption'].font.italic=False
styles['Caption'].paragraph_format.space_before=Pt(4); styles['Caption'].paragraph_format.space_after=Pt(9)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
footer.style='Normal'; footer.paragraph_format.space_after=Pt(0)

blocks=re.split(r'\n\s*\n',source.strip())
table_number=0
for block in blocks:
    if block.startswith('# '):
        lines=block.splitlines(); doc.add_paragraph(lines[0][2:],style='Title')
        if len(lines)>1: doc.add_paragraph(lines[1].removeprefix('## '),style='Subtitle')
    elif block.startswith('### '):
        lines=block.splitlines()
        for line in lines:
            if line.startswith('#### '): doc.add_paragraph(line[5:],style='Heading 2')
            elif line.startswith('### '): doc.add_paragraph(line[4:],style='Heading 1')
    elif block.startswith('#### '): doc.add_paragraph(block[5:],style='Heading 2')
    elif block.startswith('[EQUATION] '):
        p=doc.add_paragraph(block[11:]); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_together=True; p.paragraph_format.space_before=Pt(3)
        for r in p.runs: r.italic=True; r.font.size=Pt(10.5)
    elif block.startswith('[FIGURE '):
        lines=block.splitlines(); name=lines[0][8:-1]
        p=doc.add_paragraph(); p.paragraph_format.keep_with_next=True
        p.add_run().add_picture(str(ASSETS/name),width=Inches(width))
        if len(lines)>1: doc.add_paragraph(' '.join(lines[1:]),style='Caption')
    elif block.startswith('|'):
        table_number+=1
        rows=[[c.strip() for c in line.strip().strip('|').split('|')] for line in block.splitlines()]
        widths={1:[1.26,1.25,2.54,1.55],2:[1.10,1.40,1.40,.80,.70,1.20],3:[.75,1.3,1.1,1.12,1.08,1.15],4:[.81,.57,1.29,1.29,.62,.88,1.14]}[table_number]
        widths=[x*width/sum(widths) for x in widths]
        table=doc.add_table(rows=len(rows),cols=len(rows[0])); table.autofit=False
        table.alignment=WD_TABLE_ALIGNMENT.CENTER
        for col,w in zip(table.columns,widths): col.width=Inches(w)
        borders=OxmlElement('w:tblBorders')
        for edge in ['top','left','bottom','right','insideH','insideV']:
            el=OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
        table._tbl.tblPr.append(borders)
        for i,(row,vals) in enumerate(zip(table.rows,rows)):
            trpr=row._tr.get_or_add_trPr(); no=OxmlElement('w:cantSplit'); trpr.append(no)
            if i==0: trpr.append(OxmlElement('w:tblHeader'))
            for j,(cell,value,w) in enumerate(zip(row.cells,vals,widths)):
                cell.width=Inches(w); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr=cell._tc.get_or_add_tcPr(); margins=OxmlElement('w:tcMar')
                for side in ['top','left','bottom','right']:
                    el=OxmlElement('w:'+side); el.set(qn('w:w'),'85' if side in ['top','bottom'] else '75'); el.set(qn('w:type'),'dxa'); margins.append(el)
                tcpr.append(margins)
                shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'333333' if i==0 else ('F4F4F4' if i%2==0 else 'FFFFFF')); tcpr.append(shade)
                p=cell.paragraphs[0]; p.paragraph_format.line_spacing=1.05; p.paragraph_format.space_after=Pt(0)
                p.paragraph_format.keep_with_next=i<len(rows)-1
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 or table_number==1 else WD_ALIGN_PARAGRAPH.CENTER
                r=p.add_run(value); r.font.size=Pt(9.2 if table_number==4 else 9.7)
                r.bold=i==0; r.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '000000')
        doc.add_paragraph().paragraph_format.space_after=Pt(2)
    elif block.startswith('Table '):
        p=doc.add_paragraph(block,style='Caption'); p.paragraph_format.keep_with_next=True
    else:
        p=doc.add_paragraph(block.replace('\n',' '))
        if block.startswith('AI Minesweeper'):
            for r in p.runs: r.font.size=Pt(10)
            p.paragraph_format.space_after=Pt(14)
        elif block.startswith('Keywords:'):
            for r in p.runs: r.italic=True; r.font.size=Pt(10)
        elif re.match(r'^\[\d+\]',block):
            p.paragraph_format.space_after=Pt(6)
            p.paragraph_format.left_indent=Inches(.22); p.paragraph_format.first_line_indent=Inches(-.22)
            for r in p.runs: r.font.size=Pt(9.5)

# Confirm all prose blocks were included before delivering.
assert len(doc.tables)==4
assert len(doc.inline_shapes)==2
assert any('AdamW' in p.text for p in doc.paragraphs)
assert any(p.text.startswith('Minesweeper combines exact local') for p in doc.paragraphs)
assert any(p.text.startswith('Kaye established') for p in doc.paragraphs)
assert all(len(p.text)<160 for p in doc.paragraphs if p.style.name.startswith('Heading'))
all_text=' '.join(p.text for p in doc.paragraphs)
for block in blocks:
    if not block.startswith(('#','|','[FIGURE','[EQUATION')): assert block.replace('\n',' ') in all_text, block[:60]
doc.core_properties.title='Exact Inference and Selective Learning for Minesweeper'
doc.core_properties.subject='Paired evaluation and deployment aligned counterfactual adaptation'
doc.core_properties.author='AI Minesweeper Project'
doc.core_properties.keywords='Minesweeper, exact inference, counterfactual learning, paired evaluation'
doc.core_properties.comments=''
path=OUT/'minesweeper-research-paper.docx'; doc.save(path)
summary={'standard':{m:{'v44_wilson_95':wilson(w,2000),'mcnemar_exact_p':exact(b,c)} for m,w,b,c in [('beginner',1668,1,0),('intermediate',1388,5,22),('expert',690,22,55)]},'candidate':result,'verified_pairs':len(pairs),'word_count':len(source.split())}
(OUT/'paper-statistics.json').write_text(json.dumps(summary,indent=2),encoding='utf-8')
print('Created',path,'paragraphs',len(doc.paragraphs),'tables',len(doc.tables))
